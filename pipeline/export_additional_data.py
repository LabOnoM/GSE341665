#!/usr/bin/env python3
"""Export the latest Additional files and supplementary provenance tables.

The Data S8 workbook uses public openpyxl, without a proprietary artifact runtime.
"""
import argparse
import json
import shutil
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--workspace',type=Path,required=True)
    ap.add_argument('--destination',type=Path,required=True)
    ap.add_argument('--payload',type=Path,required=True)
    args=ap.parse_args()
    out=args.workspace/'scripts/JDR_reanalysis/output'
    target=args.destination/'Additional_files';target.mkdir(parents=True,exist_ok=True)
    deg=pd.read_csv(out/'DEG_table.csv',skiprows=1,float_precision='round_trip',low_memory=False)
    deg['hs']=deg['hsapiens_homolog_associated_gene_name'].astype(str).str.upper()
    strict=deg[(deg.padj<.05)&(deg.log2FoldChange.abs()>1)].sort_values('log2FoldChange')
    assert len(strict)==563
    strict.to_csv(target/'Additional_file_6.csv',index=False)
    mapping={5:'AI_gene_directions.csv',7:'module_scores.csv',8:'deconv_proportions.csv',9:'program_celltype_specificity.csv'}
    for number,name in mapping.items():shutil.copy2(out/name,target/f'Additional_file_{number}.csv')
    inhibitor=args.workspace/'00.Manuscript/Submission_02_JMM_2026-07/4_Journal_JMM/Supplementary'
    for number,name in [(2,'Supplementary_Data_S6_inhibitor_normalized_values.csv'),(3,'Supplementary_Data_S7_inhibitor_ANOVA_Dunnett.csv')]:
        shutil.copy2(inhibitor/name,target/f'Additional_file_{number}.csv')
    payload=json.loads(args.payload.read_text())
    wb=Workbook();wb.remove(wb.active)
    for index,info in enumerate(payload['sheets'],1):
        ws=wb.create_sheet(info['name'])
        for row in info['data']:ws.append(row)
        ws.freeze_panes='B2';ws.sheet_view.showGridLines=False
        for cell in ws[1]:
            cell.font=Font(name='Arial',size=10,color='FFFFFF',bold=True)
            cell.fill=PatternFill('solid',fgColor='24435C');cell.alignment=Alignment(wrap_text=True)
        ws.row_dimensions[1].height=35
        for col in ws.columns:
            key=col[0].value;letter=col[0].column_letter
            ws.column_dimensions[letter].width=24 if key=='ensembl_gene_id' else 18
            if key in info['numeric']:
                for cell in col[1:]:
                    cell.number_format='0.00E+00' if any(t in str(key) for t in ['padj','pvalue','adjusted_P','P_unpaired']) else '0.00'
        if info['name']!='Read me':
            tab=Table(displayName=f'GeneData{index}',ref=ws.dimensions)
            tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            ws.add_table(tab)
        else:
            ws.column_dimensions['A'].width=28;ws.column_dimensions['B'].width=100
            for row in ws.iter_rows(min_row=2):
                for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
                ws.row_dimensions[row[0].row].height=42
    wb.save(target/'Additional_file_4.xlsx')
    import ai_gene_provenance as provenance
    (target/'Supplementary_Tables_S1_S2.md').write_text(
        '# Supplementary Table S1\n\n'+provenance.audit_markdown()+provenance.audit_footnote()
        +'\n\n# Supplementary Table S2\n\n'+provenance.panel_markdown()+provenance.panel_footnote())


if __name__=='__main__':main()
