#!/usr/bin/env python3
"""Compare regenerated numerical evidence with the pinned manuscript inputs."""
import argparse
import json
import xml.etree.ElementTree as ET
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent


def independent_bh(values):
    p=np.asarray(values,dtype=float)
    if len(p)!=4 or not np.isfinite(p).all() or np.any((p<=0)|(p>1)):
        raise ValueError('Each BH family requires four valid nonzero probabilities')
    order=np.argsort(p);ranked=p[order]*len(p)/np.arange(1,len(p)+1)
    answer=np.empty_like(p);answer[order]=np.minimum(1,np.minimum.accumulate(ranked[::-1])[::-1])
    return answer


def check_program_results(primary, sensitivity, analysis):
    expected={'Identity':16,'Sterol':20,'Wound':14,'IFN':18}
    if primary.program.duplicated().any() or set(primary.program)!=set(expected):
        raise ValueError('Canonical program IDs must map exactly once')
    if dict(zip(primary.program,primary.n_genes))!=expected:raise ValueError('Program sizes changed')
    if not np.allclose(independent_bh(primary.ROAST_P),primary.ROAST_BH_adjusted_P,rtol=1e-12,atol=0):
        raise ValueError('Primary BH adjustment mismatch')
    if not (primary.model.eq('~ W_1 + genotype').all() and primary.contrast.eq('dcKO minus CON').all()
            and primary.residual_df.eq(3).all() and primary.rotations.eq(99999).all()
            and primary.seed.eq(20261003).all() and primary.midp.eq(False).all()):
        raise ValueError('Primary model/contrast/rotation specification changed')
    if (primary.ROAST_P<1e-5).any():raise ValueError('ROAST probability below finite-rotation resolution')
    if len(sensitivity)!=12:raise ValueError('Missing supplementary sensitivity rows')
    expected_families={
        'ROAST_unadjusted_four_programs',
        'ROAST_mixed_RUV_adjusted_four_programs',
        'ROAST_mixed_unadjusted_four_programs'}
    if set(sensitivity.adjustment_family)!=expected_families or set(sensitivity.method)!={'ROAST','ROAST_mixed'}:
        raise ValueError('Unexpected supplementary method or adjustment family')
    for family,frame in sensitivity.groupby('adjustment_family'):
        if frame.program.duplicated().any() or set(frame.program)!=set(expected):raise ValueError('Sensitivity family IDs changed')
        if not np.allclose(independent_bh(frame.raw_P),frame.BH_adjusted_P,rtol=1e-12,atol=0):
            raise ValueError('Sensitivity BH adjustment mismatch: '+family)
    provenance=json.loads((analysis/'input_provenance.json').read_text())
    if provenance['original_DESeq2_refit'] is not False or provenance['program_limma_voom_fit'] is not True:
        raise ValueError('Fit provenance is incorrect')
    if provenance['retained_genes']!=15803 or provenance['sample_count']!=6 or provenance['program_genes_unique']!=68:
        raise ValueError('Program input dimensions changed')
    qc=json.loads((analysis/'model_qc.json').read_text())
    for model,df in [('RUV_adjusted',3),('unadjusted',4)]:
        if qc[model]['residual_df']!=df or not qc[model]['full_rank'] or not qc[model]['all_weights_finite_positive']:
            raise ValueError('Invalid design/voom weights')
    members=pd.read_csv(analysis/'program_membership.csv')
    if len(members)!=68 or members.ensembl_gene_id.duplicated().any():raise ValueError('Program membership ambiguity')
    return {'file':'ROAST contract and independent BH','primary_programs':4,'sensitivity_rows':12,'passed':True}


def compare_frames(expected, actual, name, descriptive_t=False):
    if expected.shape != actual.shape or list(expected.columns) != list(actual.columns):
        raise ValueError(f'{name}: table shape or column names changed')
    maximum = 0.0
    for column in expected:
        a, b = expected[column], actual[column]
        if not np.array_equal(a.isna(), b.isna()):raise ValueError(f'{name}/{column}: missing-value mask changed')
        if pd.api.types.is_numeric_dtype(a) and pd.api.types.is_numeric_dtype(b):
            x, y = a.dropna().to_numpy(float), b.dropna().to_numpy(float)
            if not np.array_equal(x == 0, y == 0):raise ValueError(f'{name}/{column}: zero mask changed')
            tolerance = 1e-10 if descriptive_t and column == 'P_unpaired_t' else 1e-12
            if not np.allclose(x, y, rtol=tolerance, atol=0):raise ValueError(f'{name}/{column}: numerical values changed')
            if len(x):maximum=max(maximum,float(np.max(abs(x-y))))
        elif not a.fillna('').astype(str).equals(b.fillna('').astype(str)):
            raise ValueError(f'{name}/{column}: labels or text changed')
    return {'file':name,'rows':len(expected),'columns':len(expected.columns),
            'checked_cells':expected.size,'max_absolute_difference':maximum,'passed':True}


def compare_csv(expected, actual, comment=None, descriptive_t=False):
    kwargs=dict(float_precision='round_trip',low_memory=False,comment=comment)
    return compare_frames(pd.read_csv(expected,**kwargs),pd.read_csv(actual,**kwargs),
                          expected.name,descriptive_t)


def check_reference(data_root, workspace):
    source=data_root/'01_processed_data/reproduction_inputs/output'
    output=workspace/'scripts/JDR_reanalysis/output'
    return [compare_csv(source/name,output/name) for name in
            ['GSE146855_signature_cpm.csv','GSE146855_celltype_specificity.csv','scref_celltype_counts.csv']]


def check_locus(data_root, workspace):
    source=data_root/'01_processed_data/reproduction_inputs/output'
    output=workspace/'scripts/JDR_reanalysis/output'
    return [compare_csv(source/name,output/name,comment='#') for name in
            ['smad4_loxp_reads.csv','smad4_exon_coverage.csv','smad4_exon_ratio.csv',
             'smad4_junctions.csv','smad4_locus_depth.csv']]


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--data-root',type=Path,required=True)
    ap.add_argument('--work-dir',type=Path,required=True)
    ap.add_argument('--destination',type=Path,required=True)
    ap.add_argument('--mode',choices=['manuscript','supplement','reference','reference-figure','panel','locus','all'],default='all')
    args=ap.parse_args();checks=[]
    output=args.work_dir/'scripts/JDR_reanalysis/output'
    if args.mode in ('manuscript','all'):
        primary=pd.read_csv(args.destination/'Supplementary/Supplementary_Data_S3_program_tests.csv',float_precision='round_trip')
        sensitivity=pd.read_csv(args.destination/'Supplementary/Supplementary_Data_S3_program_test_sensitivity.csv',float_precision='round_trip')
        checks.append(check_program_results(primary,sensitivity,args.destination/'Program_analysis'))
        for folder,target in [('additional_files','Additional_files'),('source_csv','Supplementary'),('inhibitor_data','Inhibitor_data')]:
            for p in sorted((REPO/'results'/folder).glob('*.csv')):
                if p.name.startswith('._'):continue
                checks.append(compare_csv(p,args.destination/target/p.name,descriptive_t=True))
        a=load_workbook(REPO/'results/validation/Additional_file_2.xlsx',read_only=True,data_only=True)
        b=load_workbook(args.destination/'Additional_files/Additional_file_2.xlsx',read_only=True,data_only=True)
        if a.sheetnames != b.sheetnames:raise ValueError('Workbook sheet names changed')
        for name in a.sheetnames:
            if name=='Read me':continue
            aa=list(a[name].values);bb=list(b[name].values)
            checks.append(compare_frames(pd.DataFrame(aa[1:],columns=aa[0]),
                                         pd.DataFrame(bb[1:],columns=bb[0]),'workbook/'+name,True))
        a.close();b.close()
        genes=pd.read_csv(args.destination/'Supplementary/Supplementary_Data_S3_all_gene_statistics.csv',float_precision='round_trip')
        if len(genes)!=15803 or genes.padj.isna().sum()!=613 or (genes.padj==0).sum()!=5:
            raise ValueError('Full gene set or adjusted-P masks changed')
    if args.mode in ('reference','reference-figure','all'):
        # The figure-only mode verifies the captured object composition, whereas
        # a complete reference rebuild must also match signatures and clusters.
        if args.mode=='reference-figure':
            name='scref_celltype_counts.csv'
            checks.append(compare_csv(args.data_root/'01_processed_data/reproduction_inputs/output'/name,output/name))
        else:checks.extend(check_reference(args.data_root,args.work_dir))
    if args.mode in ('locus','all'):checks.extend(check_locus(args.data_root,args.work_dir))
    if args.mode in ('panel','all'):
        name='AI_panel_PanelApp_v4.37.csv'
        checks.append(compare_csv(args.data_root/'01_processed_data/reproduction_inputs/output'/name,output/name))
    figures={'manuscript':[f'Figure{i}' for i in range(4,8)],
             'reference':['FigureS6_scref_umap_dotplot'],'reference-figure':['FigureS6_scref_umap_dotplot'],
             'locus':['FigureS3_smad4_loxp'],'supplement':['FigureS5_workflow','FigureS3_smad4_loxp','FigureS1_breeding_scheme']}
    required=figures.get(args.mode,[])
    if args.mode=='all':required=list(dict.fromkeys(sum(figures.values(),[])))
    # Main figures have descriptive suffixes; require one nonempty match per format.
    figure_checks=[]
    for name in required:
        extensions=('pdf','png','svg') if name in [f'Figure{i}' for i in range(4,8)] else ('pdf','png')
        for extension in extensions:
            matches=list((args.destination/'Figures').glob(name+'*.'+extension))
            if len(matches)!=1 or matches[0].stat().st_size==0:raise ValueError(f'Missing or ambiguous figure {name}.{extension}')
            if extension=='svg':
                svg=ET.parse(matches[0]).getroot()
                if svg.tag!='{http://www.w3.org/2000/svg}svg' or not svg.get('viewBox'):
                    raise ValueError(f'Invalid SVG document: {matches[0]}')
                for element in svg.iter():
                    for attribute,value in element.attrib.items():
                        if attribute.rsplit('}',1)[-1]=='href' and not value.startswith(('#','data:')):
                            raise ValueError(f'SVG requires an external resource: {matches[0]}')
            figure_checks.append(str(matches[0].relative_to(args.destination)))
    report={'passed':True,'mode':args.mode,'numerical_checks':checks,'figure_files':figure_checks,
            'raw_count_model_refit':False,'original_DESeq2_refit':False,'program_limma_voom_fit':args.mode in ('manuscript','all'),
            'tolerances':'Zero/missing masks exact; nonzero values rtol=1e-12, atol=0; descriptive NNLS t tests rtol=1e-10. ROAST seed/rotations fixed; BH independently recomputed.'}
    (args.destination/'VALIDATION.json').write_text(json.dumps(report,indent=2)+'\n')
    print(f'PASS: {len(checks)} numerical tables/sheets and {len(figure_checks)} figure files',flush=True)


if __name__=='__main__':main()
